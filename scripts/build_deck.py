"""Fills the AI Builder Cup submission template with Anvil's content (editable shapes and tables, real screenshots).
  python scripts/build_deck.py "C:/Users/<you>/Downloads/Submission Template _ AI Builder Cup.pptx"
Output: submission/Anvil - AI Builder Cup Submission.pptx   (needs: pip install python-pptx; screenshots from scripts/record_demo.py)"""
import copy
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\ADHISH SEKAR\Downloads\Submission Template _ AI Builder Cup.pptx"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "demo_recording", "screenshots")
OUT = os.path.join(ROOT, "submission", "Anvil - AI Builder Cup Submission.pptx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

ACC, INK, MUTE, LINE = RGBColor(0x34, 0x51, 0xD1), RGBColor(0x10, 0x18, 0x28), RGBColor(0x66, 0x70, 0x85), RGBColor(0xD0, 0xD5, 0xDD)
SOFT, OKC, WARN, BAD, WHITE = RGBColor(0xEE, 0xF2, 0xFF), RGBColor(0x06, 0x76, 0x47), RGBColor(0xB5, 0x47, 0x08), RGBColor(0xB4, 0x23, 0x18), RGBColor(255, 255, 255)
OKBG, WARNBG, BADBG, GREY = RGBColor(0xEC, 0xFD, 0xF3), RGBColor(0xFF, 0xFA, 0xEB), RGBColor(0xFE, 0xF3, 0xF2), RGBColor(0xF2, 0xF4, 0xF7)
FONT = "Arial"

prs = Presentation(SRC)
S = list(prs.slides)


def runs(p, text, size, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, FONT
    f.color.rgb = color
    return r


def bullet(p, indent=0.16):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for t in pPr.findall(qn("a:buNone")):
        pPr.remove(t)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def tb(slide, x, y, w, h, items, size=10.5, color=INK, anchor=MSO_ANCHOR.TOP, align=None, gap=2):
    """items: list of str | (str, dict(bold,color,size,bullet,italic,space))."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    first = True
    for it in items:
        text, o = (it, {}) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align:
            p.alignment = align
        if o.get("align"):
            p.alignment = o["align"]
        p.space_after = Pt(o.get("space", gap))
        if o.get("bullet"):
            bullet(p)
        runs(p, text, o.get("size", size), o.get("bold", False), o.get("color", color), o.get("italic", False))
    return box


def box(slide, x, y, w, h, fill=SOFT, line=None, text=None, size=10, bold=False, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, radius=0.12):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if text is not None:
        lines = text if isinstance(text, list) else [text]
        for i, ln in enumerate(lines):
            t, o = (ln, {}) if isinstance(ln, str) else ln
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = o.get("align", align)
            p.space_after = Pt(o.get("space", 1))
            if o.get("bullet"):
                bullet(p)
            runs(p, t, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", False))
    return s


def arrow(slide, x1, y1, x2, y2, color=ACC, width=1.6):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def table(slide, x, y, w, colw, rows, size=9, head_fill=ACC, rowh=0.28, bold_first_col=False):
    n, m = len(rows), len(rows[0])
    gs = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(rowh * n))
    t = gs.table
    tblPr = gs._element.graphic.graphicData.tbl.tblPr
    for k in ("firstRow", "bandRow"):
        tblPr.set(k, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"             # "No style, table grid"
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.025)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = head_fill if i == 0 else (WHITE if i % 2 else GREY)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            runs(p, str(val), size, bold=(i == 0 or (bold_first_col and j == 0)), color=WHITE if i == 0 else INK)
    return gs


def pic(slide, name, x, y, w=None, h=None):
    path = os.path.join(SHOTS, name)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def clear_body(slide, keep_ids):
    """Remove nothing from the template; helper kept for clarity."""


def set_title(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top > Emu(600000) and sh.top < Emu(800000) and sh.text_frame.text.strip():
            p = sh.text_frame.paragraphs[0]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].text = text
            return


# ------------------------------------------------------------------ slide 1: team details
s = S[0]
tfx = [sh for sh in s.shapes if sh.has_text_frame and "Team Details" in sh.text_frame.text][0].text_frame
vals = {"Team name": "Team Anvil", "Team leader name": "Adhish Sekar",
        "Problem Statement": "How can enterprises let AI agents spend money on their behalf while guaranteeing that every action stays within what a human authorized? Track: Future of work and enterprise productivity."}
for p in tfx.paragraphs:
    for k, v in vals.items():
        if p.text.strip().startswith(k):
            r = p.add_run()
            r.text = " " + v
            r.font.size = Pt(15 if k != "Problem Statement" else 12)
            r.font.name = FONT
            r.font.bold = k != "Problem Statement"
            for rr in p.runs[:-1]:
                rr.font.size = Pt(15 if k != "Problem Statement" else 12)

# ------------------------------------------------------------------ slide 3: brief
s = S[2]
tb(s, 0.3, 1.2, 9.4, 0.4, [("Anvil: a mandate enforcement layer for autonomous business spending", {"bold": True, "size": 15, "color": ACC})])
tb(s, 0.3, 1.62, 9.4, 1.75, [
    ("Anvil is the control layer between AI agents and company money. Finance signs a policy: limits, approved cities, notice and trip rules. When an employee plans a trip, a Gemini booking agent compares real offers over the A2A protocol, and Anvil audits every booking against the signed policy before any money moves: approve, send to a person, or block. A blocked agent is told exactly why and corrects itself. Approved bookings get a signed receipt verified with Google's AP2 SDK, every Gemini call is traced with its input and output, and every decision is written to a tamper-evident log.", {"size": 11.5}),
    ("It starts with travel and expense. The same engine governs software licences, procurement and cloud spend.", {"size": 11.5, "color": MUTE})], gap=6)
for i, (t, sub) in enumerate([("Judges intent", "Gemini reads what a person actually approved, not just prices"), ("Enforces in code", "The verdict is applied outside the model, so text can never override it"), ("Proves every decision", "Signed receipts, AP2 chains, hash-chained log, full AI traces")]):
    box(s, 0.3 + i * 3.17, 3.5, 3.05, 0.78, fill=SOFT, line=ACC, text=[(t, {"bold": True, "size": 12, "color": ACC}), (sub, {"size": 9.5, "color": INK})])
box(s, 0.3, 4.45, 9.4, 0.5, fill=ACC, text=[("Measured live: 19 of 19 violations caught vs 10 of 19 for a rules-only check, with 0 false blocks. A second domain (software purchases) scored 8 of 8.", {"bold": True, "size": 11, "color": WHITE})])

# ------------------------------------------------------------------ slide 4: opportunities
s = S[3]
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.startswith("Opportunities"):
        tf = sh.text_frame
        tf.paragraphs[0].runs[0].text = "Opportunities: how is it different, how does it solve the problem, what is unique"
        for p in tf.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        for r in tf.paragraphs[0].runs:
            r.font.size = Pt(15)
table(s, 0.3, 1.25, 9.4, [1.9, 2.9, 4.6], [
    ["Approach", "What it checks", "What it misses"],
    ["Budget caps and allowlists", "Price, supplier, destination", "Class or refund swaps, split charges, hidden add-ons, fake approvals: caught 10 of 19"],
    ["LLM guardrails", "Harmful or off-topic text", "Whether a purchase matches one specific signed authorization"],
    ["Google AP2 on its own", "Who authorized a payment, amount and payee limits", "Semantic drift: a $590 business-class fare at an approved price passes its limits"],
    ["Anvil", "Intent vs a signed policy, AP2 limits as a second gate, self-correction", "Prototype executor is simulated"]], size=9, rowh=0.3)
tb(s, 0.3, 3.02, 4.6, 2.2, [("How it solves the problem", {"bold": True, "size": 11, "color": ACC, "space": 3}),
    ("Finance signs the rules; a person approves anything beyond them.", {"bullet": True}),
    ("Gemini audits every booking's intent: approve, escalate or block.", {"bullet": True}),
    ("Code enforces the verdict. Only approved or human-approved bookings get a receipt.", {"bullet": True}),
    ("A blocked agent is told why, corrects itself, and is audited again.", {"bullet": True}),
    ("Every decision and every Gemini call is recorded and can be re-verified.", {"bullet": True})], size=10)
tb(s, 5.1, 3.02, 4.6, 2.2, [("What is unique (USP)", {"bold": True, "size": 11, "color": ACC, "space": 3}),
    ("Semantic drift detection against a signed policy: 19/19 vs 10/19, no false blocks.", {"bullet": True}),
    ("Closed loop: block, explain, self-correct, re-audit.", {"bullet": True}),
    ("Two independent Google gates: Gemini for judgment, the AP2 SDK for limits.", {"bullet": True}),
    ("It shows its work: every offer considered, the agent's reasoning, the exact AI input and output.", {"bullet": True}),
    ("Measured, not asserted: 26 travel and 12 software scenarios, 6 adversarial.", {"bullet": True})], size=10)

# ------------------------------------------------------------------ slide 5: features
s = S[4]
cols = [("Plan and book", ["Trip planner: round trip or one way, 12 cities, real dates", "Live estimate and policy warnings before you submit", "Live Anvil Guard feed while the agent books", "Day-by-day itinerary: flight times, cabs with from/to, stay, booking refs; print or save", "Options the agent considered, and why", "Withdraw a request; notified on decisions"]),
        ("Enforce", ["Signed policy; Gemini semantic audit with a plain-language reason", "Enforcement in code, outside the model", "Company trip rules: notice, trip length", "AP2 SDK limits as a second gate", "Self-correction loop on every block", "Finance approves a whole trip in one decision", "Free text never authorizes (prompt-injection resistant)"]),
        ("Govern and prove", ["Finance edits and signs the policy: versioned, with who, when and why", "AI traces: exact input and output of every Gemini call", "Hash-chained audit log, evidence pack export", "Counterfactual receipts: what it would have cost", "Kill switch and emergency stop", "Role separation: operators cannot approve spend", "Spend by cost centre, activity export"])]
for i, (h, items) in enumerate(cols):
    x = 0.3 + i * 3.17
    box(s, x, 1.25, 3.05, 0.34, fill=ACC, text=h, size=11.5, bold=True, color=WHITE, radius=0.2)
    box(s, x, 1.62, 3.05, 3.5, fill=GREY, line=None, radius=0.04, text=[(t, {"bullet": True, "size": 11, "align": PP_ALIGN.LEFT, "space": 5}) for t in items], anchor=MSO_ANCHOR.TOP)

# ------------------------------------------------------------------ slide 6: process flow
s = S[5]
steps = ["1  Finance signs the policy", "2  Employee plans a trip", "3  Supplier agents return offers (A2A)", "4  Gemini agent picks one offer per leg", "5  Anvil audits each booking"]
for i, t in enumerate(steps):
    x = 0.25 + i * 1.95
    box(s, x, 1.3, 1.7, 0.85, fill=SOFT if i < 4 else ACC, line=ACC, text=[(t, {"bold": True, "size": 10, "color": INK if i < 4 else WHITE})] + ([("Gemini + trip rules + AP2 limits", {"size": 8, "color": WHITE})] if i == 4 else []))
    if i < 4:
        arrow(s, x + 1.7, 1.72, x + 1.95, 1.72)
outs = [("APPROVE", OKBG, OKC, "Signed receipt and AP2 chain verified. Itinerary with booking refs."), ("ESCALATE", WARNBG, WARN, "A finance manager approves or rejects the whole trip, with the reason on screen."), ("BLOCK", BADBG, BAD, "The agent is told exactly why, picks a compliant offer, and it is audited again.")]
for i, (h, bg, fg, t) in enumerate(outs):
    x = 0.6 + i * 3.1
    arrow(s, 8.6, 2.15, x + 1.4, 2.95, color=fg)
    box(s, x, 2.95, 2.8, 1.15, fill=bg, line=fg, text=[(h, {"bold": True, "size": 11, "color": fg}), (t, {"size": 9.5})])
box(s, 0.25, 4.38, 9.5, 0.62, fill=GREY, line=LINE, text=[("Every step is written to a hash-chained audit log. Every Gemini call is traced with its input and output. A kill switch pauses any agent or revokes the policy instantly.", {"size": 10, "bold": True})])

# ------------------------------------------------------------------ slide 7: role workspaces (wireframes)
s = S[6]
roles = [("Employee", ["My trips", "Policy"], ["Plan a trip: dates, round trip, purpose, cost centre", "Live estimate and policy warnings", "Trip cards: itinerary, status, options considered"]),
         ("Finance manager", ["Approvals", "My trips", "Activity", "Spend", "Policy"], ["Trip approval card: itinerary, reason, Approve all", "Spend by cost centre; emergency stop", "Policy editor and version history"]),
         ("Platform operator", ["Overview", "Agents", "Audit trail", "AI traces", "Quality", "Demo studio"], ["Anvil Guard overview, live decisions", "Pause an agent; verify or break the chain", "AI traces: exact input and output of each call"])]
for i, (name, nav, blocks) in enumerate(roles):
    x = 0.3 + i * 3.17
    box(s, x, 1.25, 3.05, 3.85, fill=WHITE, line=LINE, radius=0.03)
    box(s, x, 1.25, 3.05, 0.32, fill=INK, shape=MSO_SHAPE.RECTANGLE, text=name, size=10.5, bold=True, color=WHITE)
    box(s, x + 0.08, 1.65, 0.82, 3.35, fill=GREY, shape=MSO_SHAPE.RECTANGLE, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, text=[(n, {"size": 8, "bold": k == 0, "color": ACC if k == 0 else MUTE, "space": 5}) for k, n in enumerate(nav)])
    for k, b in enumerate(blocks):
        box(s, x + 0.98, 1.65 + k * 1.14, 1.99, 1.04, fill=SOFT if k == 0 else GREY, line=LINE, text=b, size=8.5, radius=0.06)

# ------------------------------------------------------------------ slide 8: architecture
s = S[7]
box(s, 0.3, 1.4, 2.0, 1.25, fill=GREY, line=LINE, text=[("Browser", {"bold": True, "size": 10.5}), ("Firebase Hosting: three role workspaces (employee, finance, operator)", {"size": 8.5})])
box(s, 2.95, 1.3, 3.75, 2.2, fill=SOFT, line=ACC, anchor=MSO_ANCHOR.TOP, text=[("Cloud Run: anvil (Flask API)", {"bold": True, "size": 10.5, "color": ACC})])
for k, t in enumerate(["Booking agent (A2A client)", "Anvil Guard: audit, trip rules, AP2 gate", "Policy, approvals, controls", "Evidence, traces, hash chain"]):
    box(s, 3.05 + (k % 2) * 1.8, 1.72 + (k // 2) * 0.88, 1.7, 0.78, fill=WHITE, line=LINE, text=t, size=8.5, radius=0.08)
box(s, 7.35, 1.4, 2.35, 1.25, fill=GREY, line=LINE, text=[("Cloud Run: anvil-suppliers", {"bold": True, "size": 10.5}), ("A2A agents (Agent Cards): flight, hotel, transport; 12 cities", {"size": 8.5})])
box(s, 0.7, 4.0, 2.5, 0.95, fill=OKBG, line=OKC, text=[("Firestore", {"bold": True, "size": 10.5, "color": OKC}), ("hash-chained audit log, policy versions, AI traces, controls", {"size": 8.5})])
box(s, 3.95, 4.0, 2.5, 0.95, fill=WARNBG, line=WARN, text=[("Vertex AI: Gemini 2.5 Flash", {"bold": True, "size": 10.5, "color": WARN}), ("plan, audit and correct calls, JSON output", {"size": 8.5})])
box(s, 7.0, 4.0, 2.7, 0.95, fill=GREY, line=LINE, text=[("Secret Manager + AP2 SDK", {"bold": True, "size": 10.5}), ("orchestrator key; signed mandate chains verified per booking", {"size": 8.5})])
arrow(s, 2.3, 2.0, 2.95, 2.0)
arrow(s, 6.7, 2.0, 7.35, 2.0)
arrow(s, 7.35, 2.3, 6.7, 2.3)
arrow(s, 5.3, 3.5, 5.3, 4.0)
arrow(s, 3.6, 3.5, 2.5, 4.0)
arrow(s, 6.2, 3.5, 8.0, 4.0)
arrow(s, 1.0, 2.65, 1.0, 4.0, color=OKC)
tb(s, 2.2, 1.62, 0.9, 0.3, [("HTTPS", {"size": 7.5, "color": MUTE})])
tb(s, 6.7, 1.62, 0.8, 0.3, [("A2A", {"size": 7.5, "color": MUTE})])
tb(s, 1.05, 3.1, 1.6, 0.3, [("live feed (read-only)", {"size": 7.5, "color": OKC})])

# ------------------------------------------------------------------ slide 9: technologies
s = S[8]
table(s, 0.3, 1.25, 9.4, [1.45, 3.65, 4.3], [
    ["Layer", "Technology", "Used for"],
    ["Google AI", "Gemini 2.5 Flash on Vertex AI (google-genai SDK)", "Booking agent, intent auditor, corrector: structured JSON, temperature 0"],
    ["Agent protocols", "A2A (official a2a-sdk) and AP2 V2 SDK", "Supplier discovery over Agent Cards; signed mandate chains verified per booking"],
    ["Google Cloud", "Cloud Run (2 services), Firestore, Firebase Hosting, Secret Manager, Cloud Build", "API and agents; hash-chained log, policy versions and AI traces, live feed; web app; signing key; source deploys"],
    ["Application", "Python, Flask + gunicorn, jwcrypto / cryptography", "HTTP API, streaming runs, enforcement, signed policy, SD-JWT mandates"],
    ["Frontend", "Static HTML, CSS and JavaScript (no build), Firebase JS SDK, SVG and CSS 3D", "Three role workspaces, live Anvil Guard feed, itinerary, AI traces"],
    ["Quality", "About 340 automated checks in two modes, browser tests (Playwright), A2A interop test", "Regression safety; role flows, phone layouts, real-model smoke tests"]], size=9.5, rowh=0.52, bold_first_col=True)

# ------------------------------------------------------------------ slide 10: cost
s = S[9]
table(s, 0.3, 1.25, 9.4, [3.2, 3.0, 3.2], [
    ["Item", "Estimate", "Basis"],
    ["One Gemini audit", "about $0.002", "Measured tokens x Vertex AI list price"],
    ["One five-leg trip (1 plan + 5 audits)", "about $0.015", "Same"],
    ["10,000 audited bookings per month", "about $20 of Gemini", "10,000 x $0.002"],
    ["Cloud Run", "Scales to zero; one always-on instance about $7 to $8 a month", "List-price estimate"],
    ["Firestore, Hosting, Secret Manager", "Free tier at demo scale; under $1 per 10,000 decisions", "Estimate"],
    ["Third-party services", "None: runs only on Google Cloud", "-"]], size=10, rowh=0.44, bold_first_col=True)
tb(s, 0.3, 4.5, 9.4, 0.5, [("Estimates from list prices, not billing data. A production deployment adds real supplier and payment integrations and single sign-on.", {"size": 9.5, "color": MUTE, "italic": True})])

# ------------------------------------------------------------------ slide 11: snapshots
s = S[10]
shots = [("04_guard_done.png", "Live Anvil Guard: offers, pick, audit"), ("08_trip_approved.png", "Final itinerary with booking refs"),
         ("07_approval.png", "Finance approves a whole trip"), ("12_ai_trace.png", "AI traces: exact Gemini input and output"),
         ("10_policy_history.png", "Editable, versioned, signed policy"), ("15_protected_block.png", "Blocked, self-corrected, re-audited")]
for i, (f, cap) in enumerate(shots):
    x, y = 0.3 + (i % 3) * 3.17, 1.2 + (i // 3) * 2.0
    p = pic(s, f, x, y, w=3.05)
    p.line.color.rgb = LINE
    tb(s, x, y + 1.74, 3.05, 0.3, [(cap, {"size": 8.5, "color": MUTE})])

# ------------------------------------------------------------------ slide 12: benchmarks
s = S[11]
table(s, 0.3, 1.22, 9.4, [2.6, 3.1, 3.7], [
    ["Metric", "Result", "How it was measured"],
    ["Violations caught (travel)", "19 of 19 (rules-only baseline: 10 of 19)", "26 scenarios x 3 live Gemini repeats"],
    ["False blocks", "0 of 7 legitimate actions", "Same set, including exactly-at-the-limit cases"],
    ["Stability", "26 of 26 verdicts identical across repeats", "Temperature 0, live Gemini, none served from cache"],
    ["Adversarial scenarios", "6 of 6 caught (rules-only: 4 of 6)", "Fake approvals, authority claims, prompt injection"],
    ["Second domain (software)", "8 of 8 (rules-only: 6 of 8), 0 false blocks", "12 scenarios x 3 repeats"],
    ["Decision latency", "Median about 3 s per audit (live traces: 2.8 to 3.4 s)", "Per-call telemetry, shown on the AI traces page"],
    ["Cost per audit", "About $0.002 (estimate)", "Tokens x Vertex AI list price"],
    ["Load", "8 concurrent runs in about 34 s; audit chain intact", "Live test on Cloud Run"]], size=9.5, rowh=0.38, bold_first_col=True)
tb(s, 0.3, 4.72, 9.4, 0.4, [("The scenarios were written and labelled by our team: strong evidence, not an independent benchmark.", {"size": 9.5, "color": MUTE, "italic": True})])

# ------------------------------------------------------------------ slide 13: additional details / future
s = S[12]
cols = [("Built beyond the core", ["Official A2A client and AP2 SDK in the product", "Three role workspaces with separation of duties", "Editable, signed, versioned policy", "AI traces page and itinerary printout", "Unprotected-vs-protected contrast; guided demo", "Second domain proven (software purchases)", "POST /v1/authorize for any agent"]),
        ("Next steps", ["Company single sign-on; server-enforced roles", "Real supplier and payment integrations", "Approved destinations per grade; line-manager chain", "Anvil as an SDK or MCP tool for ADK and LangChain agents", "Approvals and alerts by email or chat", "Cancel or change an approved trip; expense claims", "KMS-backed signing; per-tenant chains"]),
        ("Limits we state openly", ["The executor is simulated in the prototype", "The demo sign-in is UI-only; the API is open", "AP2 payment data are demo values", "The log is tamper-evident, not tamper-proof", "Scenarios are team-written", "Flight and hotel data are mock supplier data"])]
for i, (h, items) in enumerate(cols):
    x = 0.3 + i * 3.17
    box(s, x, 1.25, 3.05, 0.34, fill=ACC if i < 2 else WARN, text=h, size=11.5, bold=True, color=WHITE, radius=0.2)
    box(s, x, 1.62, 3.05, 3.4, fill=GREY, radius=0.04, anchor=MSO_ANCHOR.TOP, text=[(t, {"bullet": True, "size": 11, "align": PP_ALIGN.LEFT, "space": 5}) for t in items])

# ------------------------------------------------------------------ slide 14: links
s = S[13]
tfx = [sh for sh in s.shapes if sh.has_text_frame and "Provide links" in sh.text_frame.text][0].text_frame
vals = {"GitHub Public Repository": "https://github.com/Adhi1402/anvil", "Demo Video Link": "PASTE YOUR VIDEO LINK HERE after uploading", "Final Product Link": "https://dazzling-lodge-511607-m0.web.app"}
for p in tfx.paragraphs:
    for k, v in vals.items():
        if p.text.strip().startswith(k):
            r = p.add_run()
            r.text = ":  " + v
            r.font.name = FONT
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = ACC if k != "Demo Video Link" else BAD
            for rr in p.runs[:-1]:
                rr.font.size = Pt(13)
tb(s, 0.3, 2.85, 9.4, 2.2, [("How to try it (about two minutes)", {"bold": True, "size": 12, "color": ACC, "space": 4}),
    ("Open the product link and choose a person: Rahul (employee), Adhish Sekar (finance manager) or Swathy S (platform operator), or continue as a guest judge and switch roles with 'View as'.", {"bullet": True}),
    ("As Rahul: Plan a trip (leave in under 7 days to see an approval request). As Adhish: approve the whole trip, or edit the policy. As Swathy: open AI traces to see the exact Gemini input and output.", {"bullet": True}),
    ("The first request after a quiet period can take 10 to 20 seconds.", {"bullet": True, "color": MUTE}),
    ("Built on Firebase Hosting, Firestore, Cloud Run and Gemini (Vertex AI). Sign-in is a labelled demo; the spending executor is simulated.", {"bullet": True, "color": MUTE})], size=10.5, gap=4)

# ------------------------------------------------------------------ slide 15 (blank in the template): the demo in one picture
s = S[14]
tb(s, 0.3, 0.72, 9.4, 0.5, [("Three people, one governed workflow", {"size": 20})])
rows = [("Rahul, employee", "Plans a trip: dates, round trip, cost centre. Anvil Guard compares offers; a Gemini agent books; every booking is audited. He gets a day-by-day itinerary."),
        ("Adhish, finance manager", "Approves or rejects a whole trip that went beyond policy. Edits and signs the policy. Cannot approve his own requests."),
        ("Swathy, platform operator", "Runs the agents: pauses one, reads the AI traces and the tamper-evident audit trail. Cannot release money or change limits.")]
for i, (h, t) in enumerate(rows):
    y = 1.4 + i * 1.18
    box(s, 0.3, y, 2.4, 1.0, fill=ACC, text=h, size=12, bold=True, color=WHITE)
    box(s, 2.8, y, 6.9, 1.0, fill=GREY, text=t, size=11, align=PP_ALIGN.LEFT)
box(s, 0.3, 4.95 - 0.02, 9.4, 0.3, fill=WHITE, text="Anvil never lets text, a note or an agent's claim approve spending: only the signed policy and a person's decision do.", size=10, bold=True, color=ACC)

# ------------------------------------------------------------------ remove the instruction slide (2)
ids = prs.slides._sldIdLst
sld = list(ids)[1]
prs.part.drop_rel(sld.get(qn("r:id")))
ids.remove(sld)
prs.save(OUT)
print("saved", OUT)
